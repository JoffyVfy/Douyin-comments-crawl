
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .douyin import VideoRecord

__all__ = ["export_video_comments_to_excel"]


INVALID_SHEET_CHARS = re.compile(r'[\/*?:\[\]]')


def export_video_comments_to_excel(videos: Iterable[VideoRecord], output_path: str | Path) -> Path:
    """Persist the collected comments into an Excel workbook.

    Each sheet corresponds to a single video, listing comments row by row.
    """
    output_path = Path(output_path).resolve()
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet_names_in_use: set[str] = set()
    for video in videos:
        sheet_name = _make_sheet_name(video.title, sheet_names_in_use)
        worksheet = workbook.create_sheet(title=sheet_name)
        worksheet.append(["评论者ID", "评论内容", "回复内容"])
        for comment in video.comments:
            reply_text = "\n".join(reply for reply in comment.replies if reply)
            worksheet.append([
                comment.author_id,
                comment.text,
                reply_text,
            ])
        _auto_adjust_columns(worksheet)
        sheet_names_in_use.add(sheet_name)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _make_sheet_name(raw_title: str, existing: set[str]) -> str:
    base = INVALID_SHEET_CHARS.sub("_", raw_title.strip()) or "视频"
    base = base[:31]
    candidate = base
    index = 1
    while candidate in existing:
        suffix = f"_{index}"
        candidate = (base[: 31 - len(suffix)] + suffix) if len(base) + len(suffix) > 31 else base + suffix
        index += 1
    return candidate


def _auto_adjust_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            if cell.value is None:
                continue
            value = str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        adjusted = min(max_length + 2, 80)
        worksheet.column_dimensions[get_column_letter(column)].width = adjusted

